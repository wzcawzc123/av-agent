"""文档生成 Agent：按 deliverables 生成 PPT / PDF / 偏离表，产出 context.files 对应 key。"""
import os

from app.agents.base_agent import BaseAgent
from app.agents.registry import register

_SYSTEM_CN = {
    "prosound": "专业扩声", "speech": "会议发言", "display": "显示系统",
    "paperless": "无纸化会议", "control": "中控矩阵", "distributed": "分布式",
    "lighting": "灯光系统", "broadcast": "公共广播", "videoconf": "视频会议",
}

_FILE_KEYS = {"ppt": "ppt", "pdf": "pdf", "deviation": "deviation"}


def _build_ppt_fallback(slots: dict, devices: list[dict], template_path, out_path: str) -> str:
    """无 LLM 时的确定性 PPT：标题页 + 每系统一页设备要点。"""
    from pptx import Presentation

    if template_path:
        prs = Presentation(template_path)
        blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
    else:
        prs = Presentation()
        blank = prs.slide_layouts[1]

    first = prs.slides.add_slide(blank)
    if first.shapes.title is not None:
        first.shapes.title.text = f"{slots.get('scene') or '音视频'}项目方案"

    by_system: dict[str, list[dict]] = {}
    for d in devices:
        by_system.setdefault(d.get("system") or "", []).append(d)
    for code, devs in by_system.items():
        slide = prs.slides.add_slide(blank)
        if slide.shapes.title is not None:
            slide.shapes.title.text = _SYSTEM_CN.get(code, code)
        body = None
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title and shape.text_frame.text == "":
                body = shape.text_frame
                break
        if body is not None:
            for i, d in enumerate(devs[:6]):
                p = body.paragraphs[0] if i == 0 else body.add_paragraph()
                p.text = (f"{d.get('type', '')} {d.get('spec', '')} "
                          f"× {d.get('qty', 1)}{d.get('unit', '')}")
    prs.save(out_path)
    return out_path


class DocumentGenerationAgent(BaseAgent):
    name = "document_generation"
    description = "生成 PPT / PDF / 偏离表等文档交付物"

    async def analyze(self, context) -> None:
        deliverables = context.cfg.get("deliverables") or []
        self._targets = [d for d in ("ppt", "pdf", "deviation") if d in deliverables]
        if not self._targets:
            return
        if context.session is None:
            raise ValueError("缺少数据库会话（context.session）")

    async def execute(self, context) -> None:
        from app.generators.excel_generator import generate_deviation_sheet
        from app.generators.pdf_converter import convert_docx_to_pdf
        from app.generators.ppt_generator import build_ppt

        slots = context.slots
        project_dir = context.cfg.get("project_dir") or ""
        if not project_dir:
            from app.config import settings
            project_dir = f"{settings.OUTPUT_DIR}/proj_{context.project_id}"
        os.makedirs(project_dir, exist_ok=True)
        tpl_paths = dict(context.cfg.get("template_paths") or {})
        devices = context.bom or []
        targets = getattr(self, "_targets", [])

        if "ppt" in targets:
            out = os.path.join(project_dir, "方案.pptx")
            if context.provider is None:
                _build_ppt_fallback(slots, devices, tpl_paths.get("ppt"), out)
            else:
                await build_ppt(context.provider, slots, devices, tpl_paths.get("ppt"), out)
            context.files["ppt"] = out
        if "pdf" in targets:
            src = context.files.get("doc") or os.path.join(project_dir, "方案.docx")
            out = os.path.join(project_dir, "方案.pdf")
            if convert_docx_to_pdf(src, out):
                context.files["pdf"] = out
            else:
                context.errors[self.name] = "PDF 转换失败（需先有 Word，且电脑安装 LibreOffice）"
        if "deviation" in targets:
            out = os.path.join(project_dir, "偏离表.xlsx")
            generate_deviation_sheet(
                tpl_paths.get("deviation"),
                [{"requirement": slots.get("scene") or "需求", "status": "满足", "note": ""}],
                out,
            )
            context.files["deviation"] = out
        context.outputs[self.name] = {"files": dict(context.files)}

    async def validate(self, context) -> bool:
        targets = getattr(self, "_targets", [])
        if not targets:
            return True
        if any(context.files.get(_FILE_KEYS[t]) for t in targets):
            return True
        # 未产出但失败原因已写入 errors（如 PDF 依赖 Word），视为已记录不重复报错
        return bool(context.errors.get(self.name))


register(DocumentGenerationAgent.name, DocumentGenerationAgent)
