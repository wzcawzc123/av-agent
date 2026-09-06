from openpyxl import Workbook, load_workbook


def parse_design_sheet(path) -> list[dict]:
    """从设计方案清单 xlsx 解析回 BOM 行（供编辑回填）。"""
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    rows, header_idx = [], None
    for row in ws.iter_rows(values_only=True):
        cells = ["" if c is None else str(c).strip() for c in row]
        if not any(cells):
            continue
        if "产品名称" in cells and "型号" in cells:
            header_idx = {c: i for i, c in enumerate(cells)}
            continue
        if header_idx is None:
            continue
        first = cells[0]
        if first.startswith("一、") or first.startswith("二、") or first == "报价合计(元)":
            continue
        if cells[1] not in ("主要设备", "配件辅材"):
            continue
        def g(*names):
            for n in names:
                if n in header_idx and header_idx[n] < len(cells) and cells[header_idx[n]]:
                    return cells[header_idx[n]]
            return ""
        rows.append({
            "category": cells[1],
            "type": g("产品名称"),
            "spec": g("规格", "产品规格"),
            "brand": g("品牌"),
            "model": g("型号", "产品型号"),
            "qty": int(float(g("数量") or 0) or 0),
            "unit": g("单位") or "台",
            "price": float(g("单价", "单价(元)") or 0),
            "note": g("备注"),
        })
    wb.close()
    return rows


def generate_deviation_sheet(template_path, items: list[dict], out_path: str) -> str:
    if template_path:
        wb = load_workbook(template_path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(["需求", "状态", "说明"])
    for it in items:
        ws.append([it.get("requirement", ""), it.get("status", "满足"), it.get("note", "")])
    wb.save(out_path)
    return out_path


def build_deviation_sheet(out_path, header, rows, template_path=None):
    """偏离表：序号 | 货物名称 | 招标文件要求 | 投标文件实际情况 | 响应情况 | 说明"""
    wb = Workbook() if not template_path else load_workbook(template_path)
    ws = wb.active
    if not template_path:
        ws.append(["序号", "货物名称", "招标文件要求", "投标文件实际情况", "响应情况", "说明"])
    for i, r in enumerate(rows, start=1):
        ws.append([r.get("seq", i), r.get("device", ""), r.get("tender_param", ""),
                   r.get("bid_param", ""), r.get("deviation", ""), r.get("note", "")])
    wb.save(out_path)
    return out_path


def build_meeting_list(out_path, header, rows, template_path=None):
    """会议清单：项目信息 + 序号|产品名称|规格|品牌|型号|数量|单位|单价|总价|备注"""
    wb = Workbook() if not template_path else load_workbook(template_path)
    ws = wb.active
    ws.append(["", "项目名称", header.get("项目名称", ""), "", "报价日期", header.get("报价日期", "")])
    ws.append(["", "项目公司", header.get("项目公司", ""), "", "报价单位", header.get("报价单位", "")])
    ws.append(["序号", "产品名称", "产品规格", "品牌", "产品型号", "数量", "单位", "单价", "总价", "备注"])
    for i, r in enumerate(rows, start=1):
        ws.append([i, r.get("name", ""), r.get("spec", ""), r.get("brand", ""),
                   r.get("model", ""), r.get("qty", 1), r.get("unit", "台"),
                   r.get("price", 0), r.get("total", 0), r.get("note", "")])
    wb.save(out_path)
    return out_path


def build_broadcast_list(out_path, header, zones, rows, template_path=None):
    """广播清单：点位汇总（分区/明细/功率/功放）+ 设备清单"""
    wb = Workbook() if not template_path else load_workbook(template_path)
    ws = wb.active
    ws.append(["公共广播系统点位汇总"])
    ws.append(["分区", "喇叭明细", "总功率(W)", "选用功放"])
    for z in zones:
        detail = " ".join(f"{k}x{v}" for k, v in z.items()
                          if k not in ("zone", "power_w", "amplifier"))
        ws.append([z.get("zone", ""), detail, z.get("power_w", 0), z.get("amplifier", "")])
    ws.append([])
    ws.append(["序号", "产品名称", "产品型号", "数量", "单位", "备注"])
    for i, r in enumerate(rows, start=1):
        ws.append([i, r.get("name", ""), r.get("model", ""),
                   r.get("qty", 1), r.get("unit", "只"), r.get("note", "")])
    wb.save(out_path)
    return out_path


def build_led_list(out_path, header, layout, rows, template_path=None):
    """LED 清单：屏体尺寸/分辨率 + 设备清单 + 功耗/电缆"""
    wb = Workbook() if not template_path else load_workbook(template_path)
    ws = wb.active
    ws.append(["屏体尺寸", f"{layout.get('actual_w_m', 0)}m × {layout.get('actual_h_m', 0)}m"])
    ws.append(["分辨率", f"{layout.get('res_w', 0)} × {layout.get('res_h', 0)}"])
    ws.append(["序号", "产品名称", "产品型号", "数量", "单位", "备注"])
    for i, r in enumerate(rows, start=1):
        ws.append([i, r.get("name", ""), r.get("model", ""),
                   r.get("qty", 1), r.get("unit", "块"), r.get("note", "")])
    ws.append([])
    ws.append(["总功耗(KW)", layout.get("power_kw", 0), "电缆线径(mm²)", layout.get("cable_mm2", 0)])
    wb.save(out_path)
    return out_path


def build_design_sheet(out_path, header, devices, template_path=None):
    """设计方案设备清单：一、主要设备；二、配件辅材。

    列：序号 | 类别 | 产品名称 | 规格 | 品牌 | 型号 | 数量 | 单位 | 单价 | 总价 | 备注
    """
    wb = Workbook() if not template_path else load_workbook(template_path)
    ws = wb.active
    ws.append(["", "项目名称", header.get("项目名称", ""), "", "方案日期", header.get("方案日期", "")])
    ws.append(["序号", "类别", "产品名称", "规格", "品牌", "型号",
               "数量", "单位", "单价(元)", "总价(元)", "备注"])
    # 引擎/流程产出行 category 为「主设备」，BOM 编辑与招标确认行统一为「主要设备」，二者都算主设备
    main_rows = [d for d in devices if d.get("category") in ("主设备", "主要设备")]
    acc_rows = [d for d in devices if d.get("category") not in ("主设备", "主要设备")]
    seq = 0
    if main_rows:
        ws.append(["一、主要设备"] + [""] * 9)
        for d in main_rows:
            seq += 1
            price = d.get("market_price") or d.get("base_price") or 0
            qty = d.get("qty", 1)
            ws.append([seq, "主要设备", d.get("type", ""), d.get("spec", ""),
                       d.get("brand", ""), d.get("model", ""), qty, d.get("unit", "台"),
                       price, round(price * qty, 2), d.get("note", "")])
    if acc_rows:
        ws.append([])
        ws.append(["二、配件辅材"] + [""] * 9)
        for d in acc_rows:
            seq += 1
            price = d.get("market_price") or d.get("base_price") or 0
            qty = d.get("qty", 1)
            ws.append([seq, "配件辅材", d.get("type", ""), d.get("spec", ""),
                       d.get("brand", ""), d.get("model", ""), qty, d.get("unit", "只"),
                       price, round(price * qty, 2), d.get("note", "")])
    # 报价汇总
    total = sum(float(d.get("market_price") or d.get("base_price") or 0) * d.get("qty", 1)
                for d in devices)
    pending = sum(1 for d in devices if not (d.get("market_price") or d.get("base_price")))
    ws.append([])
    ws.append(["报价合计(元)", total])
    if pending:
        ws.append(["注", f"{pending} 项待询价（单价为 0，按询价结果更新产品库价格后重新生成即可）"])
    wb.save(out_path)
    return out_path
