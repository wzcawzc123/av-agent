import json

from pptx import Presentation

from app.llm.base import ChatMessage
from app.llm.prompts import PPT_PROMPT
from app.orchestrator.intent import extract_json


async def build_ppt(provider, slots: dict, devices: list[dict], template_path, out_path: str) -> str:
    user_msg = (
        f"项目：{json.dumps(slots, ensure_ascii=False)}\n"
        f"设备清单：{json.dumps(devices, ensure_ascii=False)}"
    )
    resp = await provider.chat([
        ChatMessage("system", PPT_PROMPT),
        ChatMessage("user", user_msg),
    ], temperature=0.5)
    data = extract_json(resp)
    slides = data.get("slides", [])
    if template_path:
        prs = Presentation(template_path)
        blank = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
    else:
        prs = Presentation()
        blank = prs.slide_layouts[1]
    for s in slides:
        slide = prs.slides.add_slide(blank)
        title = slide.shapes.title
        if title is not None:
            title.text = s.get("title", "")
        body = None
        for shape in slide.shapes:
            if shape.has_text_frame and shape != title and shape.text_frame.text == "":
                body = shape.text_frame
                break
        if body is not None:
            for i, b in enumerate(s.get("bullets", [])):
                p = body.paragraphs[0] if i == 0 else body.add_paragraph()
                p.text = b
    prs.save(out_path)
    return out_path
