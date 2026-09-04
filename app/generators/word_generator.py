import re

from docx import Document
from docx.shared import RGBColor


def _split_body(body: str) -> list[tuple[str, str]]:
    """把 markdown 文本拆成 (style, text)：heading/bullet/para。"""
    out = []
    for line in body.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("#### "):
            out.append(("heading", line[5:]))
        elif line.startswith("### "):
            out.append(("heading", line[4:]))
        elif line.startswith("## "):
            out.append(("heading", line[3:]))
        elif line.startswith("# "):
            out.append(("heading", line[2:]))
        elif line.startswith("- ") or line.startswith("* "):
            out.append(("bullet", line[2:]))
        elif re.match(r"^\d+[.、]", line):
            out.append(("bullet", re.sub(r"^\d+[.、]\s*", "", line)))
        elif "【图：" in line and "】" in line:
            out.append(("image_hint", line))
        else:
            out.append(("para", line))
    return out


def fill_docx_template(template_path: str | None, replacements: dict, out_path: str) -> str:
    pattern = re.compile(r"\{\{\s*(\w+)\s*\}\}")

    if template_path:
        doc = Document(template_path)
    else:
        doc = Document()
        doc.add_paragraph("{{项目名称}}")
        doc.add_paragraph("{{项目概述}}")
        doc.add_paragraph("{{方案正文}}")

    for para in list(doc.paragraphs):
        text = para.text
        if "{{方案正文}}" in text:
            body = str(replacements.get("方案正文", ""))
            if body:
                anchor = para
                for style, t in _split_body(body):
                    if style == "heading":
                        np = doc.add_paragraph(t)
                        np.style = doc.styles["Heading 2"]
                    elif style == "bullet":
                        np = doc.add_paragraph(t)
                        np.style = doc.styles["List Bullet"]
                    elif style == "image_hint":
                        np = doc.add_paragraph()
                        run = np.add_run(t)
                        run.italic = True
                        run.font.color.rgb = RGBColor(0x8A, 0x8A, 0x8A)
                    else:
                        np = doc.add_paragraph(t)
                    anchor._p.addnext(np._p)
                    anchor = np
            para._p.getparent().remove(para._p)
        else:
            new_text = pattern.sub(lambda m: str(replacements.get(m.group(1), m.group(0))), text)
            if new_text != text:
                para.text = new_text
    doc.save(out_path)
    return out_path


_SYSTEM_CN = {
    "prosound": "专业扩声", "speech": "会议发言", "display": "显示系统",
    "paperless": "无纸化会议", "control": "中控矩阵", "distributed": "分布式",
    "lighting": "灯光系统", "broadcast": "公共广播", "videoconf": "视频会议",
}


def _system_name(x) -> str:
    return _SYSTEM_CN.get(str(x), str(x))


def _overview(slots: dict) -> str:
    systems = slots.get("systems") or []
    name = slots.get("scene") or "音视频"
    parts = [f"{name}项目"]
    if slots.get("area"):
        parts.append(f"面积约{slots['area']}㎡")
    if slots.get("seats"):
        s = slots["seats"]
        parts.append(f"共{s.get('chairman', 0) + s.get('delegate', 0)}个发言席位"
                     f"（主席{s.get('chairman', 0)}个、代表{s.get('delegate', 0)}个）")
    if systems:
        parts.append("涵盖系统：" + "、".join(_system_name(x) for x in systems))
    if slots.get("brand"):
        parts.append(f"品牌要求：{slots['brand']}")
    return "；".join(parts) + "。"


async def build_doc_from_llm(provider, slots: dict, devices: list[dict],
                             template_path: str | None, out_path: str) -> str:
    from app.llm.base import ChatMessage
    from app.llm.prompts import DOC_PROMPT

    devices_txt = "\n".join(
        f"- {d['type']} {d.get('spec', '')} × {d['qty']}" for d in devices)
    user_msg = f"项目：{slots}\n设备清单：\n{devices_txt}"
    body = await provider.chat([
        ChatMessage("system", DOC_PROMPT),
        ChatMessage("user", user_msg),
    ], temperature=0.5)
    return fill_docx_template(template_path,
                              {"项目名称": slots.get("scene") or "音视频方案",
                               "项目概述": _overview(slots),
                               "方案正文": body}, out_path)
