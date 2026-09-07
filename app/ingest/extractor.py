"""文件文本提取：把任意上传文件转成 LLM 可读的文本/表格描述。

支持 xlsx/xlsm、docx、pdf、txt/md/csv。
Excel 按"单元格坐标 + 值"输出，保留表格结构语义；输出截断避免超出上下文。
"""

from __future__ import annotations

import csv
import os

MAX_TEXT_CHARS = 30_000


def extract_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm", ".xls"):
        return _extract_excel(path)
    if ext == ".docx":
        return _extract_docx(path)
    if ext == ".pdf":
        return _extract_pdf(path)
    if ext in (".txt", ".md", ".csv", ".json", ".log"):
        return _extract_plain(path)
    return f"（不支持的格式 {ext}，仅支持 Excel/Word/PDF/文本）"


def _truncate(text: str) -> str:
    return text[:MAX_TEXT_CHARS]


def _extract_excel(path: str) -> str:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return "（Excel 解析失败，文件可能损坏或加密）"
    parts = []
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        # 空 sheet 跳过
        if not any(any(c is not None and str(c).strip() for c in row) for row in rows):
            continue
        parts.append(f"【工作表：{ws.title}】")
        for ri, row in enumerate(rows[:400], start=1):
            cells = []
            for ci, cell in enumerate(row):
                if cell is None:
                    continue
                val = str(cell).strip()
                if not val:
                    continue
                cells.append(f"{chr(65 + ci)}列{ri}行={val}")
            if cells:
                parts.append(" | ".join(cells))
        if len(rows) > 400:
            parts.append(f"（表格共 {len(rows)} 行，已截断）")
    return _truncate("\n".join(parts))


def _extract_docx(path: str) -> str:
    from docx import Document

    try:
        doc = Document(path)
    except Exception:
        return "（Word 解析失败）"
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return _truncate("\n".join(parts))


def _extract_pdf(path: str) -> str:
    import pdfplumber

    try:
        with pdfplumber.open(path) as pdf:
            parts = []
            for page in pdf.pages[:20]:
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(text)
            return _truncate("\n".join(parts))
    except Exception:
        return "（PDF 解析失败）"


def _extract_plain(path: str) -> str:
    try:
        if path.endswith(".csv"):
            with open(path, encoding="utf-8-sig", errors="replace") as f:
                rows = list(csv.reader(f))
            return _truncate("\n".join(" | ".join(r) for r in rows[:500]))
        with open(path, encoding="utf-8", errors="replace") as f:
            return _truncate(f.read())
    except Exception:
        return "（文本读取失败）"
