"""BOM：设计方案清单 价格/合计/待询价 + 解析回填 + 重建 normalize"""
import json

from openpyxl import load_workbook

from app.generators.excel_generator import build_design_sheet, parse_design_sheet

ROWS = [
    {"category": "主要设备", "type": "视频会议终端", "spec": "4K 超清",
     "brand": "MAXHUB", "model": "BM65", "qty": 1, "unit": "台",
     "market_price": 8800, "note": ""},
    {"category": "主要设备", "type": "阵列麦克风", "spec": "全向拾音",
     "brand": "MAXHUB", "model": "BM-A1", "qty": 2, "unit": "个",
     "market_price": 0, "note": "待询价"},
    {"category": "配件辅材", "type": "HDMI线", "spec": "10米", "brand": "绿联",
     "model": "HD-10", "qty": 4, "unit": "条", "market_price": 120, "note": ""},
]


def test_design_sheet_price_total_and_pending(tmp_path):
    out = str(tmp_path / "清单.xlsx")
    build_design_sheet(out, {"项目名称": "测试会议室"}, ROWS)
    wb = load_workbook(out)
    ws = wb.active
    assert ws.cell(1, 3).value == "测试会议室"
    for row in ws.iter_rows(values_only=True):
        if row[0] == "报价合计(元)":
            assert row[1] == 9280
            break
    else:
        raise AssertionError("缺少报价合计行")
    found = False
    for row in ws.iter_rows(values_only=True):
        if row[0] == "注" and "待询价" in str(row[1] or ""):
            assert "1 项" in str(row[1])
            found = True
    assert found, "缺少待询价标注"


def test_parse_design_sheet_roundtrip(tmp_path):
    out = str(tmp_path / "清单.xlsx")
    build_design_sheet(out, {"项目名称": "测试会议室"}, ROWS)
    parsed = parse_design_sheet(out)
    assert len(parsed) == 3
    assert parsed[0]["model"] == "BM65"
    assert parsed[0]["price"] == 8800
    assert parsed[1]["price"] == 0
    assert parsed[1]["note"] == "待询价"
    assert parsed[2]["category"] == "配件辅材"
    assert parsed[2]["qty"] == 4


def test_rebuild_rows_normalize(tmp_path):
    rows = [
        {"category": "主要设备", "type": "功放", "brand": "itc",
         "model": "T-2350", "qty": 1, "unit": "台", "price": 3200, "note": ""},
    ]
    for r in rows:
        r.setdefault("qty", 1)
        r.setdefault("unit", "台")
        r.setdefault("note", "")
        r.setdefault("category", "主要设备")
        r["market_price"] = float(r.get("market_price") or r.get("price") or 0)
        r["base_price"] = float(r.get("base_price") or 0)
    out = str(tmp_path / "v2.xlsx")
    build_design_sheet(out, {"项目名称": "测试"}, rows)
    assert parse_design_sheet(out)[0]["price"] == 3200
    assert json.dumps(rows, ensure_ascii=False).count("market_price") == 1
