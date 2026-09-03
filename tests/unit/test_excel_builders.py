"""Excel 构建器：偏离表 / 会议清单 / 广播清单 / LED 清单"""
from openpyxl import load_workbook

from app.generators.excel_generator import (
    build_broadcast_list,
    build_deviation_sheet,
    build_led_list,
    build_meeting_list,
)


def test_build_deviation_sheet(tmp_path):
    out = tmp_path / "deviation.xlsx"
    path = build_deviation_sheet(str(out), {}, [
        {"device": "矩阵", "tender_param": "1、≥4路HDMI", "bid_model": "DS-8004",
         "bid_param": "1.支持HDMI1.4", "deviation": "无偏离", "note": ""},
    ])
    wb = load_workbook(path)
    ws = wb.active
    assert ws.cell(1, 2).value == "货物名称"
    assert ws.cell(2, 2).value == "矩阵"


def test_build_meeting_list(tmp_path):
    out = tmp_path / "meeting.xlsx"
    path = build_meeting_list(str(out), {"项目名称": "测试会议室"}, [
        {"name": "全频音箱", "spec": "8寸", "brand": "itc", "model": "TK-L208",
         "qty": 4, "unit": "只", "price": 0},
    ])
    wb = load_workbook(path)
    ws = wb.active
    assert ws.cell(1, 2).value == "项目名称"
    assert ws.cell(4, 2).value == "全频音箱"
    assert ws.cell(4, 5).value == "TK-L208"


def test_build_broadcast_list(tmp_path):
    out = tmp_path / "broadcast.xlsx"
    path = build_broadcast_list(str(out), {}, [
        {"zone": "分区1", "T-601": 12, "power_w": 180, "amplifier": "T-240"}],
        [{"name": "壁挂喇叭", "model": "T-601", "qty": 12, "unit": "只"}])
    wb = load_workbook(path)
    ws = wb.active
    assert ws.cell(3, 1).value == "分区1"
    assert ws.cell(3, 4).value == "T-240"


def test_build_led_list(tmp_path):
    out = tmp_path / "led.xlsx"
    layout = {"count_w": 24, "count_h": 16, "actual_w_m": 6.0, "actual_h_m": 4.0,
              "res_w": 1536, "res_h": 1024, "power_kw": 9.36, "cable_mm2": 247}
    path = build_led_list(str(out), {"项目名称": "LED屏"}, layout,
                          [{"name": "模组", "model": "TV-PH250-YZ", "qty": 384, "unit": "块"}])
    wb = load_workbook(path)
    ws = wb.active
    assert ws.cell(1, 2).value == "6.0m × 4.0m"
    assert ws.cell(4, 3).value == "TV-PH250-YZ"
