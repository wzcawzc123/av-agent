import json

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app.generators.word_generator import fill_docx_template
from app.generators.excel_generator import generate_deviation_sheet
from app.generators.ppt_generator import build_ppt


def _make_template(path):
    doc = Document()
    doc.add_paragraph("项目名称：{{项目名称}}")
    doc.add_paragraph("{{方案正文}}")
    doc.save(path)


def test_fill_docx_template(tmp_path):
    tpl = tmp_path / "tpl.docx"
    _make_template(tpl)
    out = tmp_path / "out.docx"
    fill_docx_template(str(tpl), {"项目名称": "100平会议室", "方案正文": "正文内容"}, str(out))
    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert any("100平会议室" in t for t in texts)
    assert any("正文内容" in t for t in texts)
    assert not any("{{" in t for t in texts)


def test_deviation_sheet_no_template(tmp_path):
    out = tmp_path / "dev.xlsx"
    generate_deviation_sheet(None, [{"requirement": "支持8欧负载", "status": "满足", "note": "功放支持"}], str(out))
    wb = load_workbook(str(out))
    ws = wb.active
    assert ws.cell(1, 1).value == "需求"
    assert ws.cell(2, 1).value == "支持8欧负载"
    assert ws.cell(2, 2).value == "满足"


class FakeProvider:
    name = "fake"

    async def chat(self, messages, temperature=0.7):
        return json.dumps({"slides": [{"title": "项目概述", "bullets": ["100平会议室"]},
                                      {"title": "配置清单", "bullets": ["音箱×2"]}]})


@pytest.mark.asyncio
async def test_build_ppt_no_template(tmp_path):
    out = tmp_path / "out.pptx"
    await build_ppt(FakeProvider(), {"scene": "会议室"}, [], None, str(out))
    prs = Presentation(str(out))
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    assert any("项目概述" in t for t in texts)


def test_fill_docx_renders_markdown_and_overview(tmp_path):
    from app.generators.word_generator import fill_docx_template

    out = tmp_path / "doc.docx"
    fill_docx_template(
        None,
        {"项目名称": "测试会议室", "项目概述": "概述内容。",
         "方案正文": "# 一、项目概况\n\n这是正文第一段。\n\n## 1.1 扩声系统\n\n- 主音箱 6 只\n- 功放 3 台\n\n第二段。\n"},
        str(out),
    )
    doc = Document(str(out))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    assert paras[0] == "测试会议室"
    assert "概述内容" in paras[1]
    assert "一、项目概况" in paras[2]
    assert "这是正文第一段" in paras[3]
    assert "1.1 扩声系统" in paras[4]
    assert "主音箱 6 只" in paras[5]
    assert "功放 3 台" in paras[6]
    assert "第二段" in paras[7]
    assert "{{" not in "\n".join(paras)
