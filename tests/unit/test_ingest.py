"""LLM 智能入库单测：文件提取、LLM 解析、入库去重。"""

import pytest

from app.ingest.classifier import (
    _clean_json,
    _to_price,
    classify_document,
    extract_knowledge,
    extract_products,
)
from app.ingest.extractor import extract_text
from app.ingest.ingestor import ingest_knowledge, ingest_products, summarize
from app.db.session import get_engine, get_session


class _FakeProvider:
    def __init__(self, resp):
        self.resp = resp
        self.calls = []

    async def chat(self, messages, temperature=0.7):
        self.calls.append([m.content for m in messages])
        return self.resp


# ---------- extractor ----------

def test_extract_xlsx(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "会议平板"
    ws.append(["产品名称", "型号", "品牌", "市场价"])
    ws.append(["会议一体机 75寸", "MP-75C", "MAXHUB", 28000])
    ws.append(["无线投屏器", "WB03", "MAXHUB", 899])
    p = tmp_path / "prods.xlsx"
    wb.save(p)
    text = extract_text(str(p))
    assert "会议平板" in text
    assert "MP-75C" in text
    assert "28000" in text


def test_extract_docx(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("会议室音视频系统方案")
    doc.add_paragraph("扩声采用壁挂音箱，声压级≥85dB")
    p = tmp_path / "doc.docx"
    doc.save(p)
    text = extract_text(str(p))
    assert "音视频系统方案" in text
    assert "85dB" in text


def test_extract_unsupported(tmp_path):
    p = tmp_path / "a.zip"
    p.write_bytes(b"x")
    assert "不支持" in extract_text(str(p))


# ---------- classifier ----------

def test_clean_json_fences_and_junk():
    assert _clean_json('{"a": 1}') == {"a": 1}
    assert _clean_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert _clean_json('说明：\n{"products": []}') == {"products": []}
    assert _clean_json("不是 json") is None


def test_to_price_variants():
    assert _to_price("¥28,000") == 28000.0
    assert _to_price("899元") == 899.0
    assert _to_price(None) == 0.0
    assert _to_price("N/A") == 0.0


@pytest.mark.asyncio
async def test_classify_document_products():
    provider = _FakeProvider('{"type":"products","title":"产品清单","summary":"含会议平板","confidence":0.95}')
    res = await classify_document(provider, "表格内容", "prods.xlsx")
    assert res["type"] == "products"
    assert res["title"] == "产品清单"


@pytest.mark.asyncio
async def test_classify_document_unknown_falls_back():
    provider = _FakeProvider('{"type":"weird"}')
    res = await classify_document(provider, "x", "a.txt")
    assert res["type"] == "other"


@pytest.mark.asyncio
async def test_extract_products_parses_and_cleans():
    provider = _FakeProvider(
        '{"products": [{"name":"会议一体机","model":"MP-75C","brand":"MAXHUB",'
        '"base_price":"20000","market_price":"¥28,000","unit":"台",'
        '"params":{"功率":"400W","尺寸":"75寸","接口":"HDMI×2"}}]}'
    )
    items = await extract_products(provider, "表格")
    assert len(items) == 1
    assert items[0]["model"] == "MP-75C"
    assert items[0]["market_price"] == 28000.0
    assert items[0]["unit"] == "台"
    assert items[0]["params"] == {"功率": "400W", "尺寸": "75寸", "接口": "HDMI×2"}


@pytest.mark.asyncio
async def test_extract_knowledge():
    provider = _FakeProvider('{"title":"会议室方案要点","doc_type":"方案案例","excerpt":"扩声采用壁挂音箱"}')
    doc = await extract_knowledge(provider, "长文本", "a.docx")
    assert doc["title"] == "会议室方案要点"
    assert doc["doc_type"] == "方案案例"


# ---------- ingestor ----------

@pytest.fixture
def engine(tmp_path):
    eng = get_engine(f"sqlite:///{tmp_path}/ingest.db")
    from app.db.models import Base
    from app.db.migrate import ensure_schema

    Base.metadata.create_all(eng)
    ensure_schema(eng)
    return eng


def test_ingest_products_dedup(engine):
    from app.db.models import Product

    items = [
        {"name": "会议一体机", "model": "MP-75C", "brand": "MAXHUB", "market_price": 28000},
        {"name": "会议一体机", "model": "MP-75C", "brand": "MAXHUB"},  # 重复型号
        {"name": "无线投屏器", "model": "WB03", "brand": "MAXHUB"},
    ]
    with get_session(engine) as s:
        added, skipped, warns = ingest_products(s, items)
        assert added == 2
        assert skipped == 1
        assert s.query(Product).count() == 2


def test_ingest_knowledge_dedup_by_title(engine):
    doc = {"title": "会议室要点", "doc_type": "资料", "excerpt": "内容"}
    with get_session(engine) as s:
        first = ingest_knowledge(s, doc)
        second = ingest_knowledge(s, doc)
        assert first > 0
        assert second == 0


def test_summarize():
    assert "新增 3 条" in summarize("products", 3, 1)
    assert "知识库" in summarize("knowledge", 1, 0, 9)
    assert "不受支持" in summarize("other", 0, 0)
