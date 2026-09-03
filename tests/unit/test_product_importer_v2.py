"""通用产品库导入器测试：多 sheet、分类段、名称继承、价格留空、重复跳过。"""
import pytest
from openpyxl import Workbook

from app.db.session import get_engine, get_session
from app.db.models import Base, Product
from app.db.product_importer import import_products_v2


def _save(wb, path):
    wb.save(path)
    return str(path)


def make_maxhub_style(tmp_path):
    """模拟 MAXHUB 清单：表头 + 分类段 + 名称继承 + 扩展列。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "会议平板"
    ws.append(["产品名称", "型号", "品牌", "产品描述", "产品图片", "尺寸（长*宽*高）", "建议视距", "适用范围"])
    ws.append(["1--会议平板"])
    ws.append(["1.0—标准款"])
    ws.append(["标准款", "EG65MZ", "MAXHUB", "65英寸超高清LCD屏", None, "1490×915×99mm", "2.6m-5.7m", "10m²-30m²"])
    ws.append([None, "EG75MZ", "MAXHUB", "75英寸超高清LCD屏", None, "1711×1039×99mm", "2.8m-6.6m", "30m²-50m²"])
    ws.append(["1.1—经典款"])
    ws.append(["经典款 V7", "CG65MA", "MAXHUB", "65英寸UHD屏", None, "1488×908×90mm", "2.6m-5.7m", "10m²-30m²"])
    ws2 = wb.create_sheet("专业拾扩音（导入）")
    ws2.append(["序号", "系统分类", "品牌", "销售型号", "产品名称", "开票名称", "描述", "产品图片"])
    ws2.append([1, "扩声系统", "MAXHUB", "MH-VS10", "10寸多功能专业音箱", "音箱", "1.系统:2-way全频音箱", None])
    ws2.append([2, "扩声系统", "MAXHUB", "MH-VS12", "12寸多功能专业音箱", "音箱", "1.系统:12寸两分频", None])
    return _save(wb, tmp_path / "maxhub.xlsx")


def make_av_style(tmp_path):
    """模拟 20241104 清单：标题行 + 表头 + 分类段 + 单价数量列。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "音视频产品"
    ws.append(["MAXHUB 音视频产品清单"])
    ws.append(["序号", "产品名称", "产品型号", "品牌", "产品描述", "单价", "数量", "单位", "总价", "图片"])
    ws.append(["一、MCU平台类"])
    ws.append([1, "视频会议MCU", "MH-MCU1000", "MAXHUB", "1.采用嵌入式操作系统", 1, 1, "台", 0, None])
    ws.append([2, "超融合MCU", "MH-MCU2000C", "MAXHUB", "1.全编全解技术", 1, 1, "台", 0, None])
    return _save(wb, tmp_path / "av.xlsx")


@pytest.fixture
def engine(tmp_path):
    e = get_engine(f"sqlite:///{tmp_path}/t.db")
    Base.metadata.create_all(e)
    return e


def test_maxhub_style_multi_sheet(tmp_path, engine):
    with get_session(engine) as s:
        r = import_products_v2(make_maxhub_style(tmp_path), s)
        assert r["inserted"] == 5 and r["updated"] == 0
        got = s.query(Product).filter_by(model="EG65MZ").first()
        assert got.name == "标准款"
        assert got.category == "会议平板/标准款"
        assert "65英寸" in got.description
        got2 = s.query(Product).filter_by(model="EG75MZ").first()
        assert got2.name == "标准款"
        assert got2.category == "会议平板/标准款"
        got3 = s.query(Product).filter_by(model="MH-VS10").first()
        assert got3.name == "10寸多功能专业音箱"
        assert got3.category == "专业拾扩音/扩声系统"
        assert got3.base_price == 0 and got3.market_price == 0


def test_av_style_with_price_columns(tmp_path, engine):
    with get_session(engine) as s:
        r = import_products_v2(make_av_style(tmp_path), s)
        assert r["inserted"] == 2
        got = s.query(Product).filter_by(model="MH-MCU1000").first()
        assert got.name == "视频会议MCU"
        assert got.category == "音视频产品/MCU平台类"
        assert got.description == "1.采用嵌入式操作系统"


def test_duplicate_model_skipped(tmp_path, engine):
    path = make_maxhub_style(tmp_path)
    with get_session(engine) as s:
        import_products_v2(path, s)
    with get_session(engine) as s:
        r = import_products_v2(path, s)
        assert r["inserted"] == 0 and r["updated"] == 0
        assert s.query(Product).count() == 5


def test_sheet_with_single_title_row(tmp_path, engine):
    wb = Workbook()
    ws = wb.active
    ws.title = "WpsReserved_CellImgList"
    ws.append(["WPS 内部图片列表"])
    path = _save(wb, tmp_path / "wps.xlsx")
    with get_session(engine) as s:
        r = import_products_v2(path, s)
        assert r["inserted"] == 0
