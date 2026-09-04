"""P2 模板闭环单测：多维匹配、BOM 回存、doc 模板选型、图片占位渲染。"""
import json

import pytest

from app.db.models import Template, ConfigTemplate
from app.db.template_store import (save_config_template, find_config_template,
                                   save_template, find_doc_template)
from app.generators.word_generator import fill_docx_template, _split_body


@pytest.fixture()
def db(tmp_path):
    from app.db.session import get_engine, get_session
    from app.db.migrate import ensure_schema
    from app.db.models import Base
    eng = get_engine(f"sqlite:///{tmp_path / 'tpl.db'}")
    Base.metadata.create_all(eng)
    ensure_schema(eng)
    with get_session(eng) as s:
        yield s
    with get_session(eng) as s:
        for c in s.query(ConfigTemplate).filter(ConfigTemplate.name.like("测试模板%")).all():
            s.delete(c)
        for t in s.query(Template).filter(Template.name.like("测试模板%")).all():
            s.delete(t)


def test_config_template_multidim_match(db):
    save_config_template(db, "测试模板-小型", 80, "会议室", {"rows": []},
                         systems=["speech", "prosound"], config_level="标准", brand="惠威")
    save_config_template(db, "测试模板-大型", 280, "多功能厅", {"rows": []},
                         systems=["speech", "prosound", "display", "paperless"],
                         config_level="豪华", brand="MAXHUB")

    # 场景+系统命中优先
    best = find_config_template(db, 260, "多功能厅", ["speech", "display"], "豪华", "MAXHUB")
    assert best and best.name == "测试模板-大型"
    # 面积近似 + 场景
    best2 = find_config_template(db, 90, "会议室", ["speech"])
    assert best2 and best2.name == "测试模板-小型"
    # 无任何匹配也应返回最近面积模板（不抛错）
    best3 = find_config_template(db, 1000)
    assert best3 is not None


def test_bom_roundtrip_via_store(db):
    rows = [{"system": "speech", "type": "主席单元", "spec": "嵌入式",
             "brand": "MAXHUB", "model": "MH-6710C", "qty": 1, "unit": "台", "note": ""}]
    c = save_config_template(db, "测试模板-BOM", 120, "会议室", {"rows": rows},
                             systems=["speech"], config_level="标准", brand="MAXHUB")
    loaded = json.loads(c.config_json)
    assert loaded["rows"][0]["model"] == "MH-6710C"
    assert c.systems != "[]"
    assert c.brand == "MAXHUB"


def test_doc_template_brand_scene_match(db):
    save_template(db, "测试模板-通用doc", "doc", "/tmp/generic.docx", meta={})
    save_template(db, "测试模板-MAXHUB会议", "doc", "/tmp/maxhub.docx",
                  meta={"scene": "会议室", "brand": "MAXHUB"})
    best = find_doc_template(db, "会议室", "MAXHUB")
    assert best and best.file_path == "/tmp/maxhub.docx"
    # 无匹配返回 None（走默认模板）
    assert find_doc_template(db, "体育馆", "JBL") is None


def test_split_body_image_hint():
    blocks = _split_body("## 一、概况\n\n【图：会议室全景照片】\n\n正文内容。")
    styles = [s for s, _ in blocks]
    assert styles == ["heading", "image_hint", "para"]


def test_fill_docx_image_hint_rendered(tmp_path):
    out = tmp_path / "doc.docx"
    fill_docx_template(
        None,
        {"项目名称": "测试", "项目概述": "概述。",
         "方案正文": "# 一、概况\n\n【图：会议室全景照片】\n\n正文。"},
        str(out),
    )
    from docx import Document
    doc = Document(str(out))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    assert any("会议室全景照片" in p for p in paras)
    # 占位行使用斜体灰色（image_hint 样式）
    for p in doc.paragraphs:
        if "【图：" in p.text:
            run = p.runs[0]
            assert run.italic is True
