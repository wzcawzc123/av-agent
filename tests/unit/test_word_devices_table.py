"""Word 方案参数表测试：方案末尾必须包含数据驱动的设备参数表格。"""

import asyncio

import pytest
from docx import Document


class _FakeProvider:
    async def chat(self, messages, temperature=0.7):
        return "# 系统设计\n会议室扩声采用专业音箱。\n\n## 扩声系统\n壁挂音箱覆盖全场。"


def _build(tmp_path):
    from app.generators.word_generator import build_doc_from_llm

    devices = [
        {"category": "主设备", "type": "专业音箱", "spec": "8寸 壁挂/吸顶",
         "brand": "MAXHUB", "model": "AV-8A", "qty": 2, "unit": "只", "note": "按覆盖计算"},
        {"category": "主设备", "type": "专业功放", "spec": "2×250W@8Ω",
         "brand": "", "model": "PA-400", "qty": 1, "unit": "台", "note": ""},
        {"category": "配件辅材", "type": "音箱线", "spec": "2×1.5mm²",
         "brand": "", "model": "", "qty": 200, "unit": "米", "note": "按实结算"},
    ]
    out = str(tmp_path / "方案.docx")
    asyncio.run(build_doc_from_llm(
        _FakeProvider(),
        {"area": 100, "scene": "会议室", "budget": "5万", "brand": "惠威"},
        devices, None, out,
    ))
    return out


def test_word_contains_devices_table(tmp_path):
    out = _build(tmp_path)
    doc = Document(out)
    # 存在设备参数表标题
    headings = [p.text for p in doc.paragraphs]
    assert any("设备清单及技术参数" in h for h in headings)
    # 表格存在且含规格/型号
    assert len(doc.tables) >= 1
    table = doc.tables[0]
    header = [c.text for c in table.rows[0].cells]
    assert "规格/参数" in header and "型号" in header and "品牌" in header
    # 行内容包含设备规格与型号（数据驱动，非 LLM 文本）
    cells_text = "\n".join(c.text for row in table.rows for c in row.cells)
    assert "8寸 壁挂/吸顶" in cells_text
    assert "AV-8A" in cells_text
    assert "2×250W@8Ω" in cells_text
    # 配件也入表
    assert "2×1.5mm²" in cells_text
